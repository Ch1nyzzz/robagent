from __future__ import annotations

import glob
import io
import os
import re
import zipfile
from typing import Optional

from agent.component_runtime.types import (
    Component, ComponentClass, ComponentContext,
    Decision, Trust,
)

_MAX_CHARS = 12000
_MEMBER_MAX_CHARS = 6000

_GAIA_GLOB = os.path.expanduser(
    "~/.cache/huggingface/hub/datasets--gaia-benchmark--GAIA"
    "/snapshots/*/2023/validation/{}"
)


def _find_cached_file(file_name: str) -> Optional[str]:
    matches = glob.glob(_GAIA_GLOB.format(file_name))
    return matches[0] if matches else None


def _strip_xml_tags(raw: str) -> str:
    text = re.sub(r"<[^>]+>", " ", raw)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text).strip()
    return text


def _zip_member_text(z: zipfile.ZipFile, name: str) -> str:
    ext = os.path.splitext(name.lower())[1]
    try:
        data = z.read(name)
        if ext in (".txt", ".csv", ".py", ".json", ".jsonld"):
            return data.decode("utf-8", errors="replace")[:_MEMBER_MAX_CHARS]
        if ext == ".xml":
            raw = data.decode("utf-8", errors="replace")
            return _strip_xml_tags(raw)[:_MEMBER_MAX_CHARS]
        if ext == ".xls":
            import xlrd  # type: ignore[import]
            wb = xlrd.open_workbook(file_contents=data)
            lines = []
            for i in range(wb.nsheets):
                ws = wb.sheet_by_index(i)
                for r in range(ws.nrows):
                    lines.append(
                        "\t".join(str(ws.cell_value(r, c)) for c in range(ws.ncols))
                    )
            return "\n".join(lines)[:_MEMBER_MAX_CHARS]
        if ext == ".xlsx":
            import openpyxl  # type: ignore[import]
            wb = openpyxl.load_workbook(io.BytesIO(data))
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append(
                        "\t".join("" if v is None else str(v) for v in row)
                    )
            return "\n".join(lines)[:_MEMBER_MAX_CHARS]
    except Exception:
        pass
    return ""


def _file_to_text(path: str, file_name: str) -> str:
    ext = os.path.splitext(file_name.lower())[1]
    try:
        if ext in (".json", ".jsonld", ".txt", ".csv", ".py"):
            with open(path, "r", errors="replace") as f:
                return f.read()[:_MAX_CHARS]
        if ext == ".xml":
            with open(path, "r", errors="replace") as f:
                raw = f.read()
            return _strip_xml_tags(raw)[:_MAX_CHARS]
        if ext == ".xlsx":
            import openpyxl  # type: ignore[import]
            wb = openpyxl.load_workbook(path)
            lines = []
            for ws in wb.worksheets:
                if ws.title:
                    lines.append(f"Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    lines.append(
                        "\t".join("" if v is None else str(v) for v in row)
                    )
            return "\n".join(lines)[:_MAX_CHARS]
        if ext == ".xls":
            import xlrd  # type: ignore[import]
            wb = xlrd.open_workbook(path)
            lines = []
            for i in range(wb.nsheets):
                ws = wb.sheet_by_index(i)
                for r in range(ws.nrows):
                    lines.append(
                        "\t".join(str(ws.cell_value(r, c)) for c in range(ws.ncols))
                    )
            return "\n".join(lines)[:_MAX_CHARS]
        if ext == ".zip":
            parts: list[str] = []
            with zipfile.ZipFile(path) as z:
                for name in z.namelist():
                    member_text = _zip_member_text(z, name)
                    if member_text:
                        parts.append(f"[File: {name}]\n{member_text}")
            return ("\n\n".join(parts))[:_MAX_CHARS]
    except Exception:
        pass
    return ""


def _matches(ctx: ComponentContext) -> bool:
    file_name = ctx.extras.get("file_name", "")
    if not file_name:
        return False
    return _find_cached_file(file_name) is not None


def _handler(ctx: ComponentContext) -> Decision:
    file_name = ctx.extras.get("file_name", "")
    path = _find_cached_file(file_name)
    if not path:
        return Decision.allow()
    content = _file_to_text(path, file_name)
    if not content:
        return Decision.allow()
    return Decision.inject_context(
        f"[Attached file content: {file_name}]\n{content}"
    )


COMPONENT = Component(
    name="gaia_file_channel",
    cls=ComponentClass.CHANNEL,
    listens="pre_prompt_build",
    matcher=_matches,
    handler=_handler,
    priority=100,
    trust=Trust(
        evidence_anchor=(
            "extras.file_name GAIA system field (set by benchmark loader from HF dataset) "
            "and HuggingFace Hub local cache path pattern "
            "'~/.cache/huggingface/hub/datasets--gaia-benchmark--GAIA/snapshots/*/2023/validation/<file_name>' "
            "— both are stable conventions documented outside the evidence traces"
        ),
        blast_radius="local",
        rollback_when=(
            "injected file content causes context to exceed model limit and truncates the task "
            "question, OR file parsing raises an unhandled exception that changes a previously "
            "correct task answer"
        ),
        out_of_evidence_probe=(
            "Task with extras.file_name pointing to an audio file (.mp3) that exists in cache: "
            "matcher fires (file found), _file_to_text returns '' (no text parser for .mp3), "
            "handler returns allow() — no injection, prompt unchanged."
        ),
        fallback=(
            "matcher returns False when file_name is empty or file absent from HF cache; "
            "handler returns allow() when file content cannot be extracted"
        ),
    ),
)
