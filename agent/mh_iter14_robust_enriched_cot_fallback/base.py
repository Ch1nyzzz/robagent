from __future__ import annotations

import re
import ssl
import urllib.request
from pathlib import Path
from typing import Any

from agent.events import EventLog, new_run_id, traces_dir
from agent.llm import chat, DEFAULT_MODEL


# ── Prompts ───────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = (
    "You are a precise problem-solving assistant. "
    "Work through the problem concisely — write only the key reasoning steps. "
    "Then output the final answer on its own line in this exact format:\n"
    "FINAL ANSWER: <value>\n"
    "where <value> is the answer only, with no extra text or punctuation."
)

RECOVERY_SYSTEM_PROMPT = (
    "You are an assistant. Output ONLY the final answer value to the question — "
    "no reasoning, no explanation, no thinking steps. "
    "Just the raw answer: a number, a short string, or a few words."
)

SEARCH_AWARE_SYSTEM_PROMPT = (
    "You are an assistant solving a benchmark task that requires looking up "
    "specific data from a named external source. "
    "You MUST search for the answer — do NOT answer directly.\n\n"
    "To find the exact data, craft a precise search query:\n"
    "- If the question names a specific website or database, add a site-specific "
    "prefix (e.g., 'site:<domain.org> <keywords>' restricts results to that source).\n"
    "- If the question provides a DOI or identifier, search for the TITLE and AUTHOR "
    "of that document rather than the identifier itself — this reaches open-access "
    "and indexed versions.\n"
    "- Include exact names, dates, and terms from the question.\n\n"
    "Respond with ONLY: SEARCH[<your targeted search query>]"
)

SEARCH_REFINE_SYSTEM_PROMPT = (
    "You are an assistant solving a benchmark task. "
    "Web search results are shown below. "
    "If the results directly answer the question, output ONLY the final answer "
    "(number, short phrase, list, etc. — no explanations). "
    "If the results are insufficient or do not contain the specific data needed, "
    "respond ONLY with: SEARCH[<a different, more targeted search query>]"
)

SEARCH_ANSWER_SYSTEM_PROMPT = (
    "You are a precise problem-solving assistant. "
    "Web search results are provided below. Use them to answer the question. "
    "Then output the final answer on its own line in this exact format:\n"
    "FINAL ANSWER: <value>\n"
    "where <value> is the answer only, with no extra text or punctuation. "
    "If the provided content clearly does not contain the specific data needed "
    "to answer the question, output EXACTLY on its own line: CANNOT ANSWER"
)

# ── Token budgets ─────────────────────────────────────────────────────────────

MAX_TOKENS = 8192
RECOVERY_MAX_TOKENS = 4096
SEARCH_MAX_TOKENS = 4096
GIANT_MAX_TOKENS = 32768

# Max characters to extract from an attached file.
FILE_MAX_CHARS = 8000

# Multi-fetch quality gate constants.
_MIN_FETCH_LEN = 300
_MAX_CSS_RATIO = 0.05
_MAX_FETCH_ATTEMPTS = 4

# ── Deterministic web-signal gate ─────────────────────────────────────────────
_WEB_SIGNAL_RE = re.compile(
    r"(?:"
    r"doi[:\s/]\s*10\.\d{4}"
    r"|https?://"
    r"|\baccording\s+to\b"
    r"|\bas\s+of\s+(?:20\d\d|\d{4})\b"
    r"|\bfrom\s+the\s+year\s+\d{4}\b"
    r"|\bthrough\s+20\d\d\b"
    r"|\bget\s+the\s+data\b"
    r"|\barxiv\b"
    r"|\byoutube\b"
    r"|\bpainting\b"
    r"|\bexhibition\b"
    r"|\bcitation\b"
    r"|\bblog\s+post\b"
    r"|\bIn\s+the\s+(?:film|movie)\b"
    r"|\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2}[,\s]+20\d\d\b"
    r")",
    re.I,
)

_SEARCH_ACTION_RE = re.compile(r"SEARCH\[([^\]]{1,400})\]")

_DOMAIN_TOKEN_RE = re.compile(
    r"\b([\w-]+\.)+(?:com|org|gov|edu|net|io|ac)\b", re.I
)

_CANNOT_ANSWER_RE = re.compile(r"CANNOT\s+ANSWER", re.I)


def _needs_web(prompt: str) -> bool:
    return bool(_WEB_SIGNAL_RE.search(prompt))


def _is_cannot_answer(content: str) -> bool:
    """Return True if the LLM explicitly signalled it cannot answer from the context."""
    return bool(content and _CANNOT_ANSWER_RE.search(content))


# ── GAIA file-reader helpers ──────────────────────────────────────────────────

# Local HuggingFace dataset cache for GAIA task attachments.
_GAIA_FILE_CACHE_BASE = (
    Path.home() / ".cache" / "huggingface" / "hub"
    / "datasets--gaia-benchmark--GAIA"
)


def _find_gaia_file(file_name: str) -> Path | None:
    """Search the local HF snapshot cache for a GAIA task attachment."""
    for snapshot_dir in _GAIA_FILE_CACHE_BASE.glob("snapshots/*/2023/validation"):
        candidate = snapshot_dir / file_name
        if candidate.exists():
            return candidate
    return None


def _strip_xml_tags(xml_text: str) -> str:
    """Strip XML/HTML tags and compress whitespace."""
    text = re.sub(r"<[^>]+>", " ", xml_text)
    return re.sub(r"\s+", " ", text).strip()


def _extract_xlsx_bytes(data: bytes, max_chars: int = FILE_MAX_CHARS) -> str:
    """Extract tabular text from an xlsx byte blob."""
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    lines: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"[Sheet: {sheet_name}]")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)[:max_chars]


def _extract_xls_bytes(data: bytes, max_chars: int = FILE_MAX_CHARS) -> str:
    """Extract tabular text from an old-format xls byte blob (via xlrd)."""
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    lines: list[str] = []
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        lines.append(f"[Sheet: {sheet_name}]")
        for i in range(ws.nrows):
            row = ws.row_values(i)
            if any(v is not None and v != "" for v in row):
                lines.append("\t".join(str(v) for v in row))
    return "\n".join(lines)[:max_chars]


def _read_file_text(file_path: Path) -> str:
    """Extract text content from a GAIA task attachment.

    Returns "" for binary formats that need vision/audio (model capability gap).
    """
    suffix = file_path.suffix.lower()

    if suffix == ".xlsx":
        try:
            return _extract_xlsx_bytes(file_path.read_bytes())
        except Exception as e:
            return f"[Error reading xlsx: {e}]"

    elif suffix == ".zip":
        import zipfile
        try:
            parts: list[str] = []
            with zipfile.ZipFile(file_path) as z:
                for name in sorted(z.namelist()):
                    inner = Path(name).suffix.lower()
                    with z.open(name) as fh:
                        raw_bytes = fh.read()
                    if inner in (".txt", ".csv", ".json", ".jsonld"):
                        parts.append(
                            f"[File: {name}]\n"
                            + raw_bytes.decode("utf-8", errors="replace")
                        )
                    elif inner == ".xml":
                        text = _strip_xml_tags(
                            raw_bytes.decode("utf-8", errors="replace")
                        )
                        parts.append(f"[File: {name}]\n{text}")
                    elif inner == ".xlsx":
                        try:
                            parts.append(
                                f"[File: {name}]\n{_extract_xlsx_bytes(raw_bytes)}"
                            )
                        except Exception as e:
                            parts.append(f"[File: {name}] (error reading: {e})")
                    elif inner == ".xls":
                        try:
                            parts.append(
                                f"[File: {name}]\n{_extract_xls_bytes(raw_bytes)}"
                            )
                        except Exception as e:
                            parts.append(f"[File: {name}] (error reading: {e})")
                    # Skip image/audio inner files silently
            return "\n\n".join(parts)[:FILE_MAX_CHARS]
        except Exception as e:
            return f"[Error reading zip: {e}]"

    elif suffix in (".jsonld", ".json", ".txt", ".csv", ".pdb"):
        try:
            return file_path.read_text(encoding="utf-8", errors="replace")[:FILE_MAX_CHARS]
        except Exception as e:
            return f"[Error reading {suffix}: {e}]"

    elif suffix == ".xml":
        try:
            raw = file_path.read_text(encoding="utf-8", errors="replace")
            return _strip_xml_tags(raw)[:FILE_MAX_CHARS]
        except Exception as e:
            return f"[Error reading xml: {e}]"

    # PDF, PPTX, PNG, JPG, MP3, etc. → model/tool capability gap, return empty.
    return ""


def _inject_file_content(task_prompt: str, file_name: str, content: str) -> str:
    """Replace the loader's [Note: not provided] placeholder with actual content."""
    placeholder = (
        f"[Note: this task references a file '{file_name}' "
        f"which is not provided in this baseline run.]"
    )
    replacement = f"[Attached file '{file_name}':\n{content}\n]"
    if placeholder in task_prompt:
        return task_prompt.replace(placeholder, replacement)
    # Fallback: append content if placeholder not found.
    return task_prompt + f"\n\n{replacement}"


# ── Web tools ─────────────────────────────────────────────────────────────────

def _web_search(query: str, log: EventLog, parent: str, max_results: int = 5) -> str:
    log.emit("tool.called", parent=parent, tool="web_search", query=query[:200])
    try:
        try:
            from ddgs import DDGS
        except ImportError:
            from duckduckgo_search import DDGS  # type: ignore[no-redef]
        with DDGS() as ddgs:
            raw = list(ddgs.text(query, max_results=max_results))
        if not raw:
            log.emit("tool.returned", parent=parent, tool="web_search", n_results=0)
            return ""
        lines: list[str] = []
        for i, r in enumerate(raw, 1):
            title = r.get("title", "")
            body = r.get("body", "")
            href = r.get("href", "")
            lines.append(f"[{i}] {title}\nURL: {href}\n{body}")
        content = "\n\n".join(lines)[:4000]
        log.emit("tool.returned", parent=parent, tool="web_search", n_results=len(raw))
        log.emit("source.opened", parent=parent, query=query[:200], source="duckduckgo")
        return content
    except Exception as e:
        log.emit("tool.failed", parent=parent, tool="web_search", error=repr(e))
        return ""


def _url_fetch(url: str, log: EventLog, parent: str, chars: int = 4000) -> str:
    """Fetch full text of a URL, stripping HTML tags. Bypasses SSL verification."""
    log.emit("tool.called", parent=parent, tool="url_fetch", url=url[:500], chars=chars)
    try:
        ctx = ssl._create_unverified_context()
        req = urllib.request.Request(
            url,
            headers={"User-Agent": "Mozilla/5.0 (compatible; research-agent/1.0)"},
        )
        with urllib.request.urlopen(req, timeout=15, context=ctx) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
        text = re.sub(r"<[^>]+>", " ", raw)
        text = re.sub(r"\s+", " ", text).strip()
        result = text[:chars]
        log.emit("tool.returned", parent=parent, tool="url_fetch", chars=len(result))
        return result
    except Exception as e:
        log.emit("tool.failed", parent=parent, tool="url_fetch", error=repr(e))
        return ""


def _passes_quality_gate(text: str) -> bool:
    """Return True if fetched page is substantive and not a CSS/nav dump."""
    if not text or len(text) < _MIN_FETCH_LEN:
        return False
    css_chars = sum(1 for c in text if c in "{};")
    return css_chars / len(text) < _MAX_CSS_RATIO


def _extract_ordered_urls(web2: str, web1: str) -> list[str]:
    """Return deduplicated URL list with search2 results prioritised."""
    seen: set[str] = set()
    result: list[str] = []
    for snippet in [web2, web1]:
        for url in re.findall(r"URL:\s*(https?://[^\s\n]+)", snippet or ""):
            if url not in seen:
                seen.add(url)
                result.append(url)
    return result


def _select_best_url(search_results: str, task_prompt: str) -> str:
    """Deterministically select the most task-relevant URL from snippet text."""
    urls = re.findall(r"URL:\s*(https?://[^\s\n]+)", search_results)
    if not urls:
        return ""
    full_domains = re.findall(
        r"\b(?:[\w-]+\.)+(?:com|org|gov|edu|net|io|ac)\b", task_prompt, re.I
    )
    for domain in full_domains:
        d = domain.lower()
        for url in urls:
            if d in url.lower():
                return url
    return urls[0]


# ── Deterministic answer extraction ──────────────────────────────────────────

def _extract_answer(content: str, finish_reason: str) -> str:
    if content:
        m = re.search(r"FINAL ANSWER:\s*(.+)", content, re.IGNORECASE)
        if m:
            candidate = m.group(1).strip().rstrip(".")
            if candidate:
                return candidate
        if finish_reason == "stop":
            lines = [ln.strip() for ln in content.strip().splitlines() if ln.strip()]
            return lines[-1] if lines else ""
    return ""


def _is_empty_length(content: str, finish_reason: str) -> bool:
    return finish_reason == "length" and not content.strip()


# ── Pipeline stages ───────────────────────────────────────────────────────────

def _run_web_search(task_prompt: str, log: EventLog, root: str) -> str:
    messages1 = [
        {"role": "system", "content": SEARCH_AWARE_SYSTEM_PROMPT},
        {"role": "user", "content": task_prompt},
    ]
    call1 = log.emit(
        "llm.requested",
        parent=root,
        messages=messages1,
        model=DEFAULT_MODEL,
        max_tokens=SEARCH_MAX_TOKENS,
        pass_="search_aware",
    )
    try:
        result1 = chat(messages=messages1, max_tokens=SEARCH_MAX_TOKENS)
    except Exception as e:
        log.emit("llm.failed", parent=call1, error=repr(e))
        return ""

    log.emit(
        "llm.responded",
        parent=call1,
        content=result1["content"],
        finish_reason=result1["finish_reason"],
        usage=result1["usage"],
        pass_="search_aware",
    )

    content1 = (result1["content"] or "").strip()

    if result1["finish_reason"] == "stop" and content1 and not _SEARCH_ACTION_RE.search(content1):
        answer = _extract_answer(content1, result1["finish_reason"])
        return answer or content1.splitlines()[-1].strip()

    action1 = _SEARCH_ACTION_RE.search(content1)
    if not action1:
        return ""

    query1 = action1.group(1).strip()
    log.emit("claim.extracted", parent=root, action="SEARCH", query=query1[:200])
    web1 = _web_search(query1, log, root)
    if not web1:
        return ""

    context2 = f"{task_prompt}\n\nWeb search results:\n{web1}"
    messages2 = [
        {"role": "system", "content": SEARCH_REFINE_SYSTEM_PROMPT},
        {"role": "user", "content": context2},
    ]
    call2 = log.emit(
        "llm.requested",
        parent=root,
        messages=messages2,
        model=DEFAULT_MODEL,
        max_tokens=SEARCH_MAX_TOKENS,
        pass_="search_refine",
    )
    try:
        result2 = chat(messages=messages2, max_tokens=SEARCH_MAX_TOKENS)
    except Exception as e:
        log.emit("llm.failed", parent=call2, error=repr(e))
        return ""

    log.emit(
        "llm.responded",
        parent=call2,
        content=result2["content"],
        finish_reason=result2["finish_reason"],
        usage=result2["usage"],
        pass_="search_refine",
    )

    content2 = (result2["content"] or "").strip()
    action2 = _SEARCH_ACTION_RE.search(content2)

    if action2 and result2["finish_reason"] == "stop":
        query2 = action2.group(1).strip()
        log.emit("claim.extracted", parent=root, action="SEARCH2", query=query2[:200])
        web2 = _web_search(query2, log, root)

        combined = (
            f"First search results:\n{web1}\n\nSecond search results:\n{web2}"
            if web2
            else web1
        )

        # Quality-gated multi-URL fetch: try search2 URLs first, then search1,
        # stopping at the first that passes the quality gate (max 4 attempts).
        ordered_urls = _extract_ordered_urls(web2 or "", web1 or "")
        page_content = ""
        best_url = ""
        for candidate_url in ordered_urls[:_MAX_FETCH_ATTEMPTS]:
            fetched = _url_fetch(candidate_url, log, root, chars=4000)
            if _passes_quality_gate(fetched):
                page_content = fetched
                best_url = candidate_url
                break

        if page_content:
            context3 = (
                f"{task_prompt}\n\nWeb search results:\n{combined}\n\n"
                f"Full page content from top result ({best_url}):\n{page_content}"
            )
        else:
            context3 = f"{task_prompt}\n\nWeb search results:\n{combined}"

        messages3 = [
            {"role": "system", "content": SEARCH_ANSWER_SYSTEM_PROMPT},
            {"role": "user", "content": context3},
        ]
        call3 = log.emit(
            "llm.requested",
            parent=root,
            messages=messages3,
            model=DEFAULT_MODEL,
            max_tokens=SEARCH_MAX_TOKENS,
            pass_="web_answer",
        )
        try:
            result3 = chat(messages=messages3, max_tokens=SEARCH_MAX_TOKENS)
        except Exception as e:
            log.emit("llm.failed", parent=call3, error=repr(e))
            return ""

        log.emit(
            "llm.responded",
            parent=call3,
            content=result3["content"],
            finish_reason=result3["finish_reason"],
            usage=result3["usage"],
            pass_="web_answer",
        )

        raw_content3 = result3["content"] or ""
        # When web_answer signals CANNOT ANSWER, the fetched web context (context3)
        # is already assembled — pass it directly to COT instead of discarding it.
        # Bare COT with no grounding hallucinates; web-enriched COT reasons over real data.
        if _is_cannot_answer(raw_content3):
            log.emit(
                "agent.cannot_answer",
                parent=root,
                reason="web_answer_cannot_answer: web-enriched COT fallback",
            )
            return _run_cot(context3, log, root)

        return _extract_answer(raw_content3, result3["finish_reason"])

    if content2 and result2["finish_reason"] == "stop":
        return _extract_answer(content2, result2["finish_reason"]) or content2.splitlines()[-1].strip()
    return ""


def _run_cot(task_prompt: str, log: EventLog, root: str) -> str:
    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": task_prompt},
    ]
    call = log.emit("llm.requested", parent=root, messages=messages, model=DEFAULT_MODEL)
    try:
        result = chat(messages=messages, max_tokens=MAX_TOKENS)
    except Exception as e:
        log.emit("llm.failed", parent=call, error=repr(e))
        return ""

    log.emit(
        "llm.responded",
        parent=call,
        content=result["content"],
        finish_reason=result["finish_reason"],
        usage=result["usage"],
    )

    answer = _extract_answer(result["content"] or "", result["finish_reason"])

    primary_empty_length = _is_empty_length(result["content"] or "", result["finish_reason"])

    if not answer:
        recovery_messages = [
            {"role": "system", "content": RECOVERY_SYSTEM_PROMPT},
            {"role": "user", "content": task_prompt},
        ]
        recovery_call = log.emit(
            "llm.requested",
            parent=root,
            messages=recovery_messages,
            model=DEFAULT_MODEL,
            recovery=True,
            recovery_max_tokens=RECOVERY_MAX_TOKENS,
        )
        try:
            recovery_result = chat(
                messages=recovery_messages,
                max_tokens=RECOVERY_MAX_TOKENS,
            )
        except Exception as e:
            log.emit("llm.failed", parent=recovery_call, error=repr(e), recovery=True)
            recovery_result = None
        else:
            log.emit(
                "llm.responded",
                parent=recovery_call,
                content=recovery_result["content"],
                finish_reason=recovery_result["finish_reason"],
                usage=recovery_result["usage"],
                recovery=True,
            )
            recovered = (recovery_result["content"] or "").strip()
            if recovered:
                answer = recovered

        if (
            not answer
            and primary_empty_length
            and recovery_result is not None
            and _is_empty_length(recovery_result["content"] or "", recovery_result["finish_reason"])
        ):
            giant_messages = [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": task_prompt},
            ]
            giant_call = log.emit(
                "llm.requested",
                parent=root,
                messages=giant_messages,
                model=DEFAULT_MODEL,
                giant_budget=True,
                giant_max_tokens=GIANT_MAX_TOKENS,
            )
            try:
                giant_result = chat(messages=giant_messages, max_tokens=GIANT_MAX_TOKENS)
            except Exception as e:
                log.emit("llm.failed", parent=giant_call, error=repr(e), giant_budget=True)
            else:
                log.emit(
                    "llm.responded",
                    parent=giant_call,
                    content=giant_result["content"],
                    finish_reason=giant_result["finish_reason"],
                    usage=giant_result["usage"],
                    giant_budget=True,
                )
                giant_content = (giant_result["content"] or "").strip()
                if giant_content:
                    answer = (
                        _extract_answer(giant_content, giant_result["finish_reason"])
                        or giant_content.splitlines()[-1].strip()
                    )

    return answer


# ── Entry point ───────────────────────────────────────────────────────────────

def run_task(
    *,
    benchmark: str,
    task_id: str,
    task_prompt: str,
    extras: dict[str, Any] | None = None,
) -> dict[str, Any]:
    run_id = new_run_id()
    log = EventLog(run_id=run_id, benchmark=benchmark, task_id=task_id, out_dir=traces_dir())
    extras = extras or {}

    root = log.emit(
        "run.started",
        question=task_prompt,
        model=DEFAULT_MODEL,
        extras=extras,
    )

    # ── Deterministic file injection ──────────────────────────────────────────
    file_name = extras.get("file_name", "")
    if file_name:
        file_path = _find_gaia_file(file_name)
        if file_path:
            file_content = _read_file_text(file_path)
            if file_content and not file_content.startswith("[Error"):
                log.emit(
                    "file.read",
                    parent=root,
                    file_name=file_name,
                    path=str(file_path),
                    content_chars=len(file_content),
                )
                task_prompt = _inject_file_content(task_prompt, file_name, file_content)
            elif not file_content:
                # Binary format (image/audio) — model capability gap.
                log.emit(
                    "agent.blocked",
                    parent=root,
                    reason=f"model_capability_gap: {file_name}",
                )
        else:
            log.emit("file.not_found", parent=root, file_name=file_name)

    # Re-evaluate web routing after potential prompt enrichment.
    web_needed = _needs_web(task_prompt)
    log.emit("task.routed", parent=root, route="WEB_SEARCH" if web_needed else "DIRECT")

    answer = ""

    if web_needed:
        answer = _run_web_search(task_prompt, log, root)

    if not answer:
        answer = _run_cot(task_prompt, log, root)

    log.emit("answer.emitted", parent=root, answer=answer)

    return {
        "run_id": run_id,
        "answer": answer,
        "error": None,
        "trace_path": str(log.path),
        "log": log,
        "root_event_id": root,
    }
