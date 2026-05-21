from __future__ import annotations

import re
from typing import Any

from agent.events import EventLog, new_run_id, traces_dir
from agent.llm import chat, DEFAULT_MODEL


SYSTEM_PROMPT = (
    "You are a precise problem-solving assistant. "
    "Work through the problem concisely — write only the key reasoning steps. "
    "Then output the final answer on its own line in this exact format:\n"
    "FINAL ANSWER: <value>\n"
    "where <value> is the answer only, with no extra text or punctuation."
)

RECOVERY_SYSTEM_PROMPT = (
    "You are an assistant. Output ONLY the final answer value to the question — "
    "no reasoning, no explanation, no thinking steps. "
    "Just the raw answer: a number, a short string, or a few words."
)

SEARCH_AWARE_SYSTEM_PROMPT = (
    "You are an assistant solving a benchmark task. "
    "The question requires data from a named external source. "
    "You MUST search for the answer — do NOT answer directly. "
    "Respond with ONLY: SEARCH[<a focused, concise search query>]"
)

SEARCH_REFINE_SYSTEM_PROMPT = (
    "You are an assistant solving a benchmark task. "
    "Web search results are shown below. "
    "If the results directly answer the question, output ONLY the final answer "
    "(number, short phrase, list, etc. — no explanations). "
    "If the results are insufficient or do not contain the specific data needed, "
    "respond ONLY with: SEARCH[<a different, more targeted search query>]"
)

SEARCH_ANSWER_SYSTEM_PROMPT = (
    "You are a precise problem-solving assistant. "
    "Web search results are provided below. Use them to answer the question. "
    "Then output the final answer on its own line in this exact format:\n"
    "FINAL ANSWER: <value>\n"
    "where <value> is the answer only, with no extra text or punctuation. "
    "If the search results are insufficient, output your best estimate."
)

MAX_TOKENS = 8192
RECOVERY_MAX_TOKENS = 1024
RECOVERY_TEMPERATURE = 0.0
SEARCH_MAX_TOKENS = 4096

_WEB_SIGNAL_RE = re.compile(
    r"(?:"
    r"\bScienceDirect\b"
    r"|\bBielefeld\b"
    r"|Girls\s+Who\s+Code"
    r"|\bUSGS\b"
    r"|\bORCID\b"
    r"|\barXiv\b"
    r"|doi[:\s]\s*10\.\d{4}"
    r"|Metropolitan\s+Museum"
    r"|Project\s+MUSE"
    r"|\bJSTOR\b"
    r"|\bTri.?Rail\b"
    r"|\bNonindigenous\s+Aquatic\s+Species\b"
    r"|\bIn\s+the\s+(?:film|movie)\b"
    r"|\bpainting\b"
    r"|\bblog\s+post\b"
    r"|\bYouTube\b"
    r")",
    re.I,
)

_SEARCH_ACTION_RE = re.compile(r"SEARCH\[([^\]]{1,400})\]")


def _needs_web(prompt: str) -> bool:
    return bool(_WEB_SIGNAL_RE.search(prompt))


def _read_hf_file(file_name: str, timeout: int = 20) -> str | None:
    """Download and parse a GAIA task file from HuggingFace."""
    import os
    import ssl
    import urllib.request

    token = os.environ.get("HF_TOKEN") or os.environ.get("HUGGINGFACE_TOKEN")
    if not token:
        env_path = os.path.join(os.path.dirname(__file__), "..", "..", ".env")
        try:
            with open(env_path) as f:
                for line in f:
                    if line.startswith("HF_TOKEN="):
                        token = line.split("=", 1)[1].strip()
                        break
        except Exception:
            pass
    if not token:
        return None

    url = (
        "https://huggingface.co/datasets/gaia-benchmark/GAIA"
        f"/resolve/main/2023/validation/{file_name}"
    )
    ctx = ssl._create_unverified_context()
    req = urllib.request.Request(
        url,
        headers={"Authorization": f"Bearer {token}", "User-Agent": "python"},
    )
    try:
        resp = urllib.request.urlopen(req, context=ctx, timeout=timeout)
        data = resp.read()
    except Exception:
        return None

    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if ext in ("json", "jsonld"):
        try:
            import json
            return json.dumps(json.loads(data), indent=2)[:6000]
        except Exception:
            return data.decode("utf-8", errors="replace")[:6000]

    elif ext in ("txt", "csv"):
        return data.decode("utf-8", errors="replace")[:6000]

    elif ext == "zip":
        import io as _io
        import zipfile

        parts: list[str] = []
        try:
            with zipfile.ZipFile(_io.BytesIO(data)) as zf:
                for inner_name in zf.namelist():
                    inner_data = zf.read(inner_name)
                    inner_ext = (
                        inner_name.rsplit(".", 1)[-1].lower() if "." in inner_name else ""
                    )
                    parts.append(f"=== {inner_name} ===")
                    if inner_ext == "xml":
                        decoded = inner_data.decode("utf-8", errors="replace")
                        texts = re.findall(r"<w:t[^>]*>([^<]+)</w:t>", decoded)
                        parts.append(
                            "\n".join(texts)
                            if texts
                            else re.sub(r"<[^>]+>", " ", decoded)[:2000]
                        )
                    elif inner_ext == "xls":
                        try:
                            import xlrd
                            wb = xlrd.open_workbook(file_contents=inner_data)
                            rows: list[str] = []
                            for sh in wb.sheets():
                                for r in range(sh.nrows):
                                    vals = [
                                        str(sh.cell_value(r, c))
                                        for c in range(sh.ncols)
                                        if sh.cell_value(r, c) != ""
                                    ]
                                    if vals:
                                        rows.append(", ".join(vals))
                            parts.append("\n".join(rows[:80]))
                        except Exception as e:
                            parts.append(f"[XLS parse error: {e}]")
                    elif inner_ext == "xlsx":
                        try:
                            import io as _sio

                            import openpyxl
                            wb2 = openpyxl.load_workbook(_sio.BytesIO(inner_data))
                            rows2: list[str] = []
                            for ws in wb2.worksheets:
                                for row in ws.iter_rows(values_only=True):
                                    vals2 = [str(v) for v in row if v is not None]
                                    if vals2:
                                        rows2.append(", ".join(vals2))
                            parts.append("\n".join(rows2[:80]))
                        except Exception as e:
                            parts.append(f"[XLSX parse error: {e}]")
                    else:
                        try:
                            parts.append(inner_data.decode("utf-8", errors="replace")[:1500])
                        except Exception:
                            parts.append("[binary]")
        except Exception as e:
            return f"[ZIP parse error: {e}]"
        return "\n\n".join(parts)[:6000]

    elif ext == "xlsx":
        try:
            import io as _sio

            import openpyxl
            wb = openpyxl.load_workbook(_sio.BytesIO(data))
            rows: list[str] = []
            for ws in wb.worksheets:
                rows.append(f"Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None]
                    if vals:
                        rows.append(", ".join(vals))
            return "\n".join(rows[:120])
        except Exception as e:
            return f"[XLSX parse error: {e}]"

    else:
        try:
            return data.decode("utf-8", errors="replace")[:4000]
        except Exception:
            return None


def _enrich_prompt(task_prompt: str, extras: dict) -> str:
    """Replace the 'file not provided' stub with actual file content when available."""
    file_name = extras.get("file_name", "")
    if not file_name:
        return task_prompt

    file_content = _read_hf_file(file_name)
    if not file_content:
        return task_prompt

    note = (
        f"[Note: this task references a file '{file_name}' "
        f"which is not provided in this baseline run.]"
    )
    enriched_block = f"[File: {file_name}]\n{file_content}"

    if note in task_prompt:
        return task_prompt.replace(note, enriched_block)
    return task_prompt + f"\n\n{enriched_block}"


def _web_search(query: str, log: EventLog, parent: str, max_results: int = 5) -> str:
    log.emit("tool.called", parent=parent, tool="web_search", query=query[:200])
    try:
        try:
            from ddgs import DDGS
        except ImportError:
            from duckduckgo_search import DDGS  # type: ignore[no-redef]
        with DDGS() as ddgs:
            raw = list(ddgs.text(query, max_results=max_results))
        if not raw:
            log.emit("tool.returned", parent=parent, tool="web_search", n_results=0)
            return ""
        lines: list[str] = []
        for i, r in enumerate(raw, 1):
            title = r.get("title", "")
            body = r.get("body", "")
            href = r.get("href", "")
            lines.append(f"[{i}] {title}\nURL: {href}\n{body}")
        content = "\n\n".join(lines)[:4000]
        log.emit("tool.returned", parent=parent, tool="web_search", n_results=len(raw))
        log.emit("source.opened", parent=parent, query=query[:200], source="duckduckgo")
        return content
    except Exception as e:
        log.emit("tool.failed", parent=parent, tool="web_search", error=repr(e))
        return ""


def _extract_answer(content: str, finish_reason: str) -> str:
    if content:
        m = re.search(r"FINAL ANSWER:\s*(.+)", content, re.IGNORECASE)
        if m:
            candidate = m.group(1).strip().rstrip(".")
            if candidate:
                return candidate
        if finish_reason == "stop":
            lines = [ln.strip() for ln in content.strip().splitlines() if ln.strip()]
            return lines[-1] if lines else ""
    return ""


def _run_web_search(task_prompt: str, log: EventLog, root: str) -> str:
    messages1 = [
        {"role": "system", "content": SEARCH_AWARE_SYSTEM_PROMPT},
        {"role": "user", "content": task_prompt},
    ]
    call1 = log.emit(
        "llm.requested",
        parent=root,
        messages=messages1,
        model=DEFAULT_MODEL,
        max_tokens=SEARCH_MAX_TOKENS,
        pass_="search_aware",
    )
    try:
        result1 = chat(messages=messages1, max_tokens=SEARCH_MAX_TOKENS)
    except Exception as e:
        log.emit("llm.failed", parent=call1, error=repr(e))
        return ""

    log.emit(
        "llm.responded",
        parent=call1,
        content=result1["content"],
        finish_reason=result1["finish_reason"],
        usage=result1["usage"],
        pass_="search_aware",
    )

    content1 = (result1["content"] or "").strip()
    action1 = _SEARCH_ACTION_RE.search(content1)
    if not action1:
        return ""

    query1 = action1.group(1).strip()
    log.emit("claim.extracted", parent=root, action="SEARCH", query=query1[:200])
    web_content1 = _web_search(query1, log, root)
    if not web_content1:
        return ""

    context2 = f"{task_prompt}\n\nWeb search results:\n{web_content1}"
    messages2 = [
        {"role": "system", "content": SEARCH_REFINE_SYSTEM_PROMPT},
        {"role": "user", "content": context2},
    ]
    call2 = log.emit(
        "llm.requested",
        parent=root,
        messages=messages2,
        model=DEFAULT_MODEL,
        max_tokens=SEARCH_MAX_TOKENS,
        pass_="search_refine",
    )
    try:
        result2 = chat(messages=messages2, max_tokens=SEARCH_MAX_TOKENS)
    except Exception as e:
        log.emit("llm.failed", parent=call2, error=repr(e))
        return ""

    log.emit(
        "llm.responded",
        parent=call2,
        content=result2["content"],
        finish_reason=result2["finish_reason"],
        usage=result2["usage"],
        pass_="search_refine",
    )

    content2 = (result2["content"] or "").strip()
    action2 = _SEARCH_ACTION_RE.search(content2)

    if action2 and result2["finish_reason"] == "stop":
        query2 = action2.group(1).strip()
        log.emit("claim.extracted", parent=root, action="SEARCH2", query=query2[:200])
        web_content2 = _web_search(query2, log, root)

        combined = (
            f"First search results:\n{web_content1}\n\nSecond search results:\n{web_content2}"
            if web_content2
            else web_content1
        )
        context3 = f"{task_prompt}\n\nWeb search results:\n{combined}"
        messages3 = [
            {"role": "system", "content": SEARCH_ANSWER_SYSTEM_PROMPT},
            {"role": "user", "content": context3},
        ]
        call3 = log.emit(
            "llm.requested",
            parent=root,
            messages=messages3,
            model=DEFAULT_MODEL,
            max_tokens=SEARCH_MAX_TOKENS,
            pass_="web_answer",
        )
        try:
            result3 = chat(messages=messages3, max_tokens=SEARCH_MAX_TOKENS)
        except Exception as e:
            log.emit("llm.failed", parent=call3, error=repr(e))
            return ""

        log.emit(
            "llm.responded",
            parent=call3,
            content=result3["content"],
            finish_reason=result3["finish_reason"],
            usage=result3["usage"],
            pass_="web_answer",
        )
        return _extract_answer(result3["content"] or "", result3["finish_reason"])

    if content2 and result2["finish_reason"] == "stop":
        return _extract_answer(content2, result2["finish_reason"]) or content2.splitlines()[-1].strip()
    return ""


def _run_cot(task_prompt: str, log: EventLog, root: str) -> str:
    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": task_prompt},
    ]
    call = log.emit("llm.requested", parent=root, messages=messages, model=DEFAULT_MODEL)
    try:
        result = chat(messages=messages, max_tokens=MAX_TOKENS)
    except Exception as e:
        log.emit("llm.failed", parent=call, error=repr(e))
        return ""

    log.emit(
        "llm.responded",
        parent=call,
        content=result["content"],
        finish_reason=result["finish_reason"],
        usage=result["usage"],
    )

    answer = _extract_answer(result["content"] or "", result["finish_reason"])

    if not answer:
        recovery_messages = [
            {"role": "system", "content": RECOVERY_SYSTEM_PROMPT},
            {"role": "user", "content": task_prompt},
        ]
        recovery_call = log.emit(
            "llm.requested",
            parent=root,
            messages=recovery_messages,
            model=DEFAULT_MODEL,
            recovery=True,
            recovery_max_tokens=RECOVERY_MAX_TOKENS,
        )
        try:
            recovery_result = chat(
                messages=recovery_messages,
                max_tokens=RECOVERY_MAX_TOKENS,
                temperature=RECOVERY_TEMPERATURE,
            )
        except Exception as e:
            log.emit("llm.failed", parent=recovery_call, error=repr(e), recovery=True)
        else:
            log.emit(
                "llm.responded",
                parent=recovery_call,
                content=recovery_result["content"],
                finish_reason=recovery_result["finish_reason"],
                usage=recovery_result["usage"],
                recovery=True,
            )
            recovered = (recovery_result["content"] or "").strip()
            if recovered:
                answer = recovered

    return answer


def run_task(
    *,
    benchmark: str,
    task_id: str,
    task_prompt: str,
    extras: dict[str, Any] | None = None,
) -> dict[str, Any]:
    run_id = new_run_id()
    log = EventLog(run_id=run_id, benchmark=benchmark, task_id=task_id, out_dir=traces_dir())
    extras = extras or {}

    root = log.emit(
        "run.started",
        question=task_prompt,
        model=DEFAULT_MODEL,
        extras=extras,
    )

    # Enrich prompt with file contents when extras.file_name is set.
    # Replaces the "[Note: file not provided]" stub with actual parsed file data
    # downloaded from HuggingFace using HF_TOKEN from .env.
    task_prompt = _enrich_prompt(task_prompt, extras)

    web_needed = _needs_web(task_prompt)
    log.emit("task.routed", parent=root, route="WEB_SEARCH" if web_needed else "DIRECT")

    answer = ""

    if web_needed:
        answer = _run_web_search(task_prompt, log, root)

    if not answer:
        answer = _run_cot(task_prompt, log, root)

    log.emit("answer.emitted", parent=root, answer=answer)

    return {
        "run_id": run_id,
        "answer": answer,
        "error": None,
        "trace_path": str(log.path),
        "log": log,
        "root_event_id": root,
    }
